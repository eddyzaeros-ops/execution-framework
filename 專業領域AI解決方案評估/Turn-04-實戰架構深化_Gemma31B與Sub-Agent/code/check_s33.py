import sys
from pptx import Presentation

sys.stdout.reconfigure(encoding='utf-8')
prs = Presentation(r"d:\JavaDO\執行框架\專業領域AI解決方案評估報告.pptx")
print("=== Slide 33 content ===")
s33 = prs.slides[32]
for shape in s33.shapes:
    if shape.has_text_frame:
        print(shape.text_frame.text[:100])
    if shape.has_table:
        for r in shape.table.rows:
            print([c.text.strip().replace("\n", " ") for c in r.cells])
