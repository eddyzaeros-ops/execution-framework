import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

pptx_path = r"d:\JavaDO\執行框架\專業領域AI解決方案評估報告.pptx"
prs = Presentation(pptx_path)
s1 = prs.slides[0]

# 找出重複的 shape 7
shape_7 = s1.shapes[7]
sp = shape_7._element
sp.getparent().remove(sp)

# 取得留下來的 shape (原本的 shape 8，現在 index 為 7)
shape_main = s1.shapes[7]

# 重新設定清晰的字型與排版
COLOR_BODY = RGBColor(0x4A, 0x55, 0x68)

for p in shape_main.text_frame.paragraphs:
    p.line_spacing = 1.3
    for r in p.runs:
        r.font.name = "Microsoft JhengHei"
        r.font.size = Pt(14)
        r.font.color.rgb = COLOR_BODY

prs.save(pptx_path)
print("Successfully fixed duplicate textbox and set clean Microsoft JhengHei font on Slide 1")
