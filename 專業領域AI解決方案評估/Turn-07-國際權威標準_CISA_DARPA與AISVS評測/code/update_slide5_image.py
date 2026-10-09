import sys
from pptx import Presentation

pptx_path = r"d:\JavaDO\執行框架\專業領域AI解決方案評估報告.pptx"
prs = Presentation(pptx_path)

# 檢視第 5 頁的 shapes
s5 = prs.slides[4]
print(f"Slide 5 shapes count: {len(s5.shapes)}")
for idx, shape in enumerate(s5.shapes):
    print(f"[{idx}] {shape.name} (type={shape.shape_type})")

# 重新載入 doc.png 替換原有圖片
# 刪除原有圖片 shape (通常是索引 0)
pic_elem = s5.shapes[0]._element
pic_elem.getparent().remove(pic_elem)

# 加入新的高解析度 doc.png (滿版 13.333 x 7.5)
from pptx.util import Inches
s5.shapes.add_picture(r"d:\JavaDO\執行框架\doc.png", Inches(0), Inches(0), width=Inches(13.333), height=Inches(7.5))

prs.save(pptx_path)
print("Successfully replaced doc.png in Slide 5 with enlarged text version!")
